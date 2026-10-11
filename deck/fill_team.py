"""Fill the team placeholders that the official template ships with.

The template's own shapes carry `[Name]`, `[ ]` and
`[Team Name] · [Contact email]`. Editing those by hand in a slide editor
re-serialises the whole package and drifts every chrome shape by a few
hundred EMU, which fails the byte-identity stage of `verify_deck.py`.
Filling them here touches only the targeted text runs, so the 33 chrome
shapes stay exactly as the template wrote them.

Run after `build_deck.py` and before `verify_deck.py`:

    deck/.venv/bin/python deck/build_deck.py
    deck/.venv/bin/python deck/fill_team.py
    deck/.venv/bin/python deck/verify_deck.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

DECK_DIR = Path(__file__).resolve().parent
PPTX = DECK_DIR / "Tech4Food_Round1_SaafScan.pptx"

# Names, exactly as the team wants them printed on the submission.
TEAM_MEMBERS = [
    "Siddhesh Kadlag (Team Lead)",
    "Harshal Pawar",
    "Prachi Kale",
]
# row index -> (text the template ships, text to write). Row 0 is the
# "Team Details" header and row 4 (Theme) is already correct in the template,
# so neither is touched.
TEAM_TABLE = {
    1: ("[ ]", "IRIS"),
    2: ("[ ]", "Siddhesh Kadlag"),
    3: ("[ ]", "Sapkal College of Engineering, Nashik"),
    5: ("Both: app + test strip", "Both: Software (app) + test strip"),
}
TEAM_FOOTER = "IRIS \u00b7 captainlevi252@gmail.com"

# The shapes these values live in, pinned by name so a template update fails
# loudly instead of silently writing into the wrong box.
SHAPE_MEMBERS = "Google Shape;29;p3"
SHAPE_TABLE = "Google Shape;30;p3"
SHAPE_FOOTER = "Google Shape;110;p11"


class FillError(RuntimeError):
    """A target shape was missing or did not hold the expected placeholder."""


def set_paragraph_text(paragraph, text: str) -> None:
    """Replace a paragraph's text, keeping the formatting of its first run."""
    runs = paragraph.runs
    if not runs:
        paragraph.text = text
        return
    runs[0].text = text
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)


def shape_by_name(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise FillError(f"shape {name!r} not found on slide")


def table_of(shape):
    if shape.shape_type != MSO_SHAPE_TYPE.TABLE:
        raise FillError(f"{shape.name!r} is not a table")
    return shape.table


def fill_members(slide) -> None:
    shape = shape_by_name(slide, SHAPE_MEMBERS)
    body = shape.text_frame
    if len(body.paragraphs) != len(TEAM_MEMBERS):
        raise FillError(
            f"{SHAPE_MEMBERS} has {len(body.paragraphs)} paragraphs, expected {len(TEAM_MEMBERS)}"
        )
    for paragraph, text in zip(body.paragraphs, TEAM_MEMBERS, strict=True):
        set_paragraph_text(paragraph, text)


def fill_table(slide) -> None:
    table = table_of(shape_by_name(slide, SHAPE_TABLE))
    for row_index, (template_text, text) in TEAM_TABLE.items():
        label = table.cell(row_index, 0).text_frame.text.strip()
        current = table.cell(row_index, 1).text_frame.text.strip()
        if current not in {template_text, text}:
            raise FillError(
                f"{SHAPE_TABLE} row {row_index} held {current!r}, expected {template_text!r}"
            )
        set_paragraph_text(table.cell(row_index, 1).text_frame.paragraphs[0], text)
        if table.cell(row_index, 1).text_frame.text.strip() != text:
            raise FillError(f"row {row_index} ({label!r}) did not take the value")


def fill_footer(slide) -> None:
    shape = shape_by_name(slide, SHAPE_FOOTER)
    set_paragraph_text(shape.text_frame.paragraphs[0], TEAM_FOOTER)


def main() -> int:
    if not PPTX.exists():
        print(f"missing {PPTX} - run build_deck.py first", file=sys.stderr)
        return 1

    prs = Presentation(PPTX)
    if len(prs.slides) != 9:
        raise FillError(f"deck has {len(prs.slides)} slides, expected 9")

    cover = prs.slides[0]
    fill_members(cover)
    fill_table(cover)
    fill_footer(prs.slides[8])

    prs.save(PPTX)
    print(f"filled team details in {PPTX.name}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except FillError as exc:
        print(f"fill_team: {exc}", file=sys.stderr)
        raise SystemExit(1) from None
